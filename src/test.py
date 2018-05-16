#!/usr/bin/python
# coding:utf-8

from openpyxl import load_workbook
import os
import re


def set_key(xlsx_file_name):
    wb = load_workbook(xlsx_file_name)
    for sheetname in wb.sheetnames:
        sheet = wb[sheetname]
        if sheet['A1'].value is not None:
            sheet['A2'].value+='=key'
            print(sheet['A2'].value)
    wb.save(xlsx_file_name)


def set_filename(xlsx_file_name):
    wb = load_workbook(xlsx_file_name)
    for sheetname in wb.sheetnames:
        sheet = wb[sheetname]
        if str(sheet['A2'].value).startswith('watergun'):
            c = sheet['A2'].value[len('watergun'):]
            filename = c.split(':')[0]
            c = c[len(filename)+1:]
            print(filename, c)
            sheet['A1'].value = filenameset_key
            sheet['A2'].value = c
    # wb.save(set_key)


def main():
    d = '../configuration/'
    files = [os.path.join(d, filename)
             for filename in os.listdir(d)]
    for filename in files:
        set_key(filename)


if __name__ == '__main__':
    main()
